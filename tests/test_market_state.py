"""Market state: percentile maths, condition branches, narrative anti-drift, fail-safe.

No qlib and no openpyxl needed — every test drives the data layer through synthetic
state dicts, so a broken qlib store cannot mask a logic regression.
"""
import numpy as np
import pandas as pd
import pytest

from qlib_lab import market_state as ms


# ------------------------------------------------------------------ percentiles

def test_trailing_pct_matches_vol_desk_formula():
    """Must equal (window <= current).mean(), NOT rolling().rank(pct=True)."""
    rng = np.random.default_rng(0)
    s = pd.Series(rng.normal(size=400))
    got = ms.trailing_pct(s, window=252)
    want = float((s.iloc[-252:] <= s.iloc[-1]).mean())
    assert got.iloc[-1] == pytest.approx(want)


def test_trailing_pct_endpoints():
    s = pd.Series(list(range(100)))
    pct = ms.trailing_pct(s, window=60)
    assert pct.iloc[-1] == pytest.approx(1.0)      # last value is the max
    falling = pd.Series(list(range(100))[::-1])
    assert ms.trailing_pct(falling, window=60).iloc[-1] == pytest.approx(1 / 60)


def test_trailing_pct_respects_min_periods():
    assert ms.trailing_pct(pd.Series(range(30)), window=252).isna().all()


# ----------------------------------------------------------------------- state

def _state(vrp_pct=0.50, vix_idx=0.50, term=0.80, netliq_z=0.0, skew_pct=0.5,
           vix_pct=0.5, absorption=0.40, crowding=None, chg4w=0.0):
    cx = crowding if crowding is not None else pd.DataFrame([
        {"asset": "SPY", "idx52": 0.50, "z156": 0.1, "chg4w": chg4w,
         "side": "neutral", "corr_to_crowded": np.nan, "same_bet": False},
    ])
    return {
        "vol": {"vrp_pct": vrp_pct, "vix": 16.0, "term_ratio": term,
                "structure": "contango" if term < 1 else "backwardation",
                "skew_pct": skew_pct, "vix_pct": vix_pct, "vvix": 90.0,
                "skew": 140.0, "rv20": 12.0, "vrp": 50.0},
        "cot_idx": {"VIX": vix_idx},
        "liq": {"netliq_chg13w_z": netliq_z, "netliq_chg13w": -50.0,
                "trend": "falling", "netliq": 5800.0},
        "funding": {"funding_1d": 0.0001},
        "struct": {"absorption": absorption, "top1_share": 0.2, "window": 60,
                   "n_assets": 13, "corr": pd.DataFrame()},
        "crowding": cx,
    }


def _fired(state):
    return {c["id"] for c in ms.conditions(state) if c["fired"]}


def _cluster(n=3, corr=0.5):
    return pd.DataFrame([
        {"asset": f"A{i}", "idx52": 0.95, "z156": 2.0, "chg4w": 0.0,
         "side": "crowded long", "corr_to_crowded": corr, "same_bet": True}
        for i in range(n)
    ])


# ------------------------------------------------------------------ conditions

def test_short_vol_crowded_while_cheap_fires_only_on_both_legs():
    assert "SHORT_VOL_CROWDED_WHILE_CHEAP" in _fired(_state(vrp_pct=0.10, vix_idx=0.10))
    assert "SHORT_VOL_CROWDED_WHILE_CHEAP" not in _fired(_state(vrp_pct=0.10, vix_idx=0.60))
    assert "SHORT_VOL_CROWDED_WHILE_CHEAP" not in _fired(_state(vrp_pct=0.60, vix_idx=0.10))


def test_crowded_cluster_needs_three_correlated_names():
    assert "CROWDED_CLUSTER" in _fired(_state(crowding=_cluster(3, 0.5)))
    assert "CROWDED_CLUSTER" not in _fired(_state(crowding=_cluster(2, 0.5)))
    # three crowded names that do NOT move together are three bets, not one
    assert "CROWDED_CLUSTER" not in _fired(_state(crowding=_cluster(3, 0.05)))


def test_term_backwardation_boundary():
    assert "TERM_BACKWARDATION" in _fired(_state(term=1.0))
    assert "TERM_BACKWARDATION" not in _fired(_state(term=0.999))


def test_liquidity_shock_and_vol_rich():
    assert "LIQUIDITY_SHOCK" in _fired(_state(netliq_z=-2.5))
    assert "LIQUIDITY_SHOCK" not in _fired(_state(netliq_z=-1.5))
    assert "VOL_RICH" in _fired(_state(vrp_pct=0.90))
    assert "VOL_RICH" not in _fired(_state(vrp_pct=0.80))


def test_skew_vs_vix_divergence():
    assert "SKEW_VS_VIX_DIVERGENCE" in _fired(_state(skew_pct=0.90, vix_pct=0.20))
    assert "SKEW_VS_VIX_DIVERGENCE" not in _fired(_state(skew_pct=0.90, vix_pct=0.50))


def test_positioning_unwind_and_absorption():
    assert "POSITIONING_UNWIND" in _fired(_state(chg4w=-0.08))
    assert "POSITIONING_UNWIND" not in _fired(_state(chg4w=-0.01))
    assert "ABSORPTION_SPIKE" in _fired(_state(absorption=0.75))
    assert "ABSORPTION_SPIKE" not in _fired(_state(absorption=0.70))


def test_every_condition_is_evaluated_and_shaped():
    out = ms.conditions(_state())
    assert len(out) == len(ms.CONDITIONS)
    assert {c["id"] for c in out} == {c["id"] for c in ms.CONDITIONS}
    for c in out:
        assert isinstance(c["fired"], bool)
        assert c["detail"] and c["rule"] and c["sentence"]


# ------------------------------------------------------------------- fail-safe

def test_conditions_neutral_on_missing_inputs():
    """Hard rule 3: missing input -> neutral, never an exception, never a fire."""
    blank = {"vol": {}, "cot_idx": {}, "liq": {}, "funding": {},
             "struct": {}, "crowding": pd.DataFrame()}
    out = ms.conditions(blank)
    assert not any(c["fired"] for c in out)
    assert all("unavailable" in c["detail"] for c in out)


def test_conditions_survive_a_broken_rule(monkeypatch):
    def boom(_):
        raise ValueError("synthetic")
    patched = [dict(ms.CONDITIONS[0], fn=boom)] + list(ms.CONDITIONS[1:])
    monkeypatch.setattr(ms, "CONDITIONS", patched)
    out = ms.conditions(_state())
    assert out[0]["fired"] is False
    assert "evaluation error" in out[0]["detail"]
    assert len(out) == len(patched)


def test_nan_inputs_do_not_fire():
    assert _fired(_state(vrp_pct=np.nan, vix_idx=np.nan, term=np.nan,
                         netliq_z=np.nan, absorption=np.nan)) == set()


# ------------------------------------------------------------------- narrative

def test_narrative_mentions_exactly_the_fired_conditions():
    """Anti-drift guarantee: prose is assembled from the same evaluation as the
    flags, so it can never claim something the flags below it do not."""
    state = _state(vrp_pct=0.10, vix_idx=0.10, absorption=0.75)
    evaluated = ms.conditions(state)
    text = ms.narrative(state, evaluated)
    for c in evaluated:
        if c["fired"]:
            assert c["sentence"] in text
        else:
            assert c["sentence"] not in text


def test_narrative_states_when_nothing_fires():
    state = _state()
    text = ms.narrative(state, ms.conditions(state))
    assert "No registered condition is firing today." in text


def test_narrative_always_carries_the_disclaimer():
    state = _state()
    assert ms.DISCLAIMER in ms.narrative(state, ms.conditions(state))


def test_narrative_reports_the_vrp_band():
    for pct, word in ((0.10, "cheap"), (0.50, "middling"), (0.90, "rich")):
        state = _state(vrp_pct=pct)
        assert word in ms.narrative(state, ms.conditions(state))


# ------------------------------------------------------------------- crowding

def test_crowding_marks_correlated_crowded_names_as_one_bet():
    pos = pd.DataFrame([
        {"asset": "GOLD", "effective_date": pd.Timestamp("2026-08-18"),
         "net_pct": 0.5, "z156": 1.2, "idx52": 0.99, "chg4w": 0.06},
        {"asset": "COPPER", "effective_date": pd.Timestamp("2026-08-18"),
         "net_pct": 0.3, "z156": 1.5, "idx52": 0.94, "chg4w": 0.01},
        {"asset": "TLT", "effective_date": pd.Timestamp("2026-08-18"),
         "net_pct": -0.2, "z156": -0.6, "idx52": 0.50, "chg4w": 0.0},
    ])
    corr = pd.DataFrame([[1.0, 0.6, 0.0], [0.6, 1.0, 0.0], [0.0, 0.0, 1.0]],
                        index=["GOLD", "COPPER", "TLT"],
                        columns=["GOLD", "COPPER", "TLT"])
    cx = ms.crowding_x_structure(pos, {"corr": corr})

    gold = cx.set_index("asset").loc["GOLD"]
    assert gold["side"] == "crowded long"
    assert gold["corr_to_crowded"] == pytest.approx(0.6)   # vs COPPER only, not self
    assert bool(gold["same_bet"]) is True
    assert cx.set_index("asset").loc["TLT"]["side"] == "neutral"


def test_crowding_lone_crowded_name_has_no_peer_correlation():
    pos = pd.DataFrame([
        {"asset": "GOLD", "effective_date": pd.Timestamp("2026-08-18"),
         "net_pct": 0.5, "z156": 1.2, "idx52": 0.99, "chg4w": 0.0},
        {"asset": "TLT", "effective_date": pd.Timestamp("2026-08-18"),
         "net_pct": 0.0, "z156": 0.0, "idx52": 0.50, "chg4w": 0.0},
    ])
    corr = pd.DataFrame([[1.0, 0.9], [0.9, 1.0]],
                        index=["GOLD", "TLT"], columns=["GOLD", "TLT"])
    cx = ms.crowding_x_structure(pos, {"corr": corr})
    gold = cx.set_index("asset").loc["GOLD"]
    assert pd.isna(gold["corr_to_crowded"])   # no OTHER crowded name to compare to
    assert bool(gold["same_bet"]) is False


def test_crowding_empty_positioning_is_empty_not_an_error():
    assert ms.crowding_x_structure(pd.DataFrame(), {"corr": pd.DataFrame()}).empty


# ------------------------------------------------------------------ registration

def test_registration_note_admits_it_is_not_blind():
    """The sheet must never let a backfilled column read as a track record."""
    assert "NOT BLIND PRE-REGISTRATION" in ms.CONDITIONS_NOTE
    assert ms.CONDITIONS_REGISTERED in ms.CONDITIONS_NOTE


def test_history_readers_are_failsafe_on_missing_files(monkeypatch, tmp_path):
    monkeypatch.setattr(ms.cot, "COT_PARQUET", tmp_path / "nope.parquet")
    monkeypatch.setattr(ms.funding, "FUND_PARQUET", tmp_path / "nope2.parquet")
    monkeypatch.setattr(ms.fred_liquidity, "daily_frame", lambda: None)
    assert ms.positioning_history().empty
    assert ms.funding_history().empty
    assert ms.liquidity_history().empty


# --------------------------------------------------- term structure staleness

def test_term_backwardation_refuses_stale_legs():
    """VIX9D and VIX3M stopped sharing sessions in this store on 2026-07-17. A ratio
    built from two different dates is not a term structure and must not fire."""
    s = _state(term=1.20)
    s["vol"]["term_stale_days"] = 4
    fired = {c["id"] for c in ms.conditions(s) if c["fired"]}
    assert "TERM_BACKWARDATION" not in fired
    detail = next(c["detail"] for c in ms.conditions(s)
                  if c["id"] == "TERM_BACKWARDATION")
    assert "stale" in detail and "different sessions" in detail


def test_term_backwardation_fires_when_legs_are_current():
    s = _state(term=1.20)
    s["vol"]["term_stale_days"] = 0
    assert "TERM_BACKWARDATION" in {c["id"] for c in ms.conditions(s) if c["fired"]}


def test_term_backwardation_absent_staleness_key_still_evaluates():
    """Older callers/synthetic states without the key must not silently stop firing."""
    s = _state(term=1.20)
    s["vol"].pop("term_stale_days", None)
    assert "TERM_BACKWARDATION" in {c["id"] for c in ms.conditions(s) if c["fired"]}


# --------------------------------------------------------------- fire log

def _fake_state(as_of, fired_ids):
    return {"as_of": as_of, "generated_at": f"{as_of}T00:00:00+00:00",
            "conditions": [{"id": i, "fired": True} for i in fired_ids]}


def _read(path):
    import json
    return [json.loads(l) for l in path.read_text().splitlines() if l.strip()]


def test_fire_log_is_idempotent_per_session_date(tmp_path):
    """One row per as_of. A second run the same day REPLACES, never duplicates -
    otherwise every fire count in the sheet inflates on a manual rebuild."""
    p = tmp_path / "log.jsonl"
    ms.append_fire_log(_fake_state("2026-08-21", ["A"]), p)
    ms.append_fire_log(_fake_state("2026-08-21", ["A", "B"]), p)
    ms.append_fire_log(_fake_state("2026-08-21", ["A", "B"]), p)
    rows = _read(p)
    assert len(rows) == 1
    assert rows[0]["fired"] == ["A", "B"]      # latest run wins


def test_fire_log_accumulates_distinct_dates_in_order(tmp_path):
    p = tmp_path / "log.jsonl"
    for d in ("2026-08-21", "2026-08-19", "2026-08-20"):
        ms.append_fire_log(_fake_state(d, ["A"]), p)
    assert [r["as_of"] for r in _read(p)] == ["2026-08-19", "2026-08-20", "2026-08-21"]


def test_fire_log_survives_a_torn_line(tmp_path):
    p = tmp_path / "log.jsonl"
    ms.append_fire_log(_fake_state("2026-08-19", ["A"]), p)
    with p.open("a", encoding="utf-8") as fh:
        fh.write('{"as_of": "2026-08-2\n')          # truncated write
    ms.append_fire_log(_fake_state("2026-08-20", ["B"]), p)
    rows = _read(p)
    assert [r["as_of"] for r in rows] == ["2026-08-19", "2026-08-20"]


def test_fire_log_never_raises_on_bad_path(tmp_path):
    ms.append_fire_log(_fake_state("2026-08-21", ["A"]),
                       tmp_path / "no" / "such" / "dir" / "x.jsonl")


# ------------------------------------------------- current-session exclusion

def test_vol_history_drops_the_live_partial_bar(monkeypatch):
    """Yahoo serves a live intraday bar for VIX while CBOE publishes only after the
    close; a midday run must not report a half-formed session as market state."""
    idx = pd.bdate_range("2024-01-01", periods=400)
    px = pd.DataFrame({t: np.linspace(10, 20, len(idx)) for t in ms.VOL_TICKERS},
                      index=idx)
    # append a bar dated 'today' with only SPY/VIX populated, as Yahoo would intraday
    today = pd.Timestamp.now().normalize()
    px.loc[today] = np.nan
    px.loc[today, ["SPY", "VIX"]] = [999.0, 99.0]
    monkeypatch.setattr(ms, "close_panel", lambda *a, **k: px)

    out = ms.vol_history(start="2024-01-01")
    assert out.index.max() < today
    assert 999.0 not in out["spy"].values


def test_vol_history_empty_when_only_todays_bar_exists(monkeypatch):
    today = pd.Timestamp.now().normalize()
    px = pd.DataFrame({t: [1.0] for t in ms.VOL_TICKERS}, index=[today])
    monkeypatch.setattr(ms, "close_panel", lambda *a, **k: px)
    assert ms.vol_history().empty


# ------------------------------------------------------- locked-workbook fallback

def test_build_workbook_falls_back_when_target_is_locked(tmp_path, monkeypatch):
    """Excel takes an exclusive lock on an open workbook. At 11:45 the user having
    it open is the NORMAL case, so a lock must degrade to a sidecar, not lose the run."""
    from qlib_lab import market_state_xlsx as msx

    target = tmp_path / "market_state.xlsx"
    target.write_text("existing")

    real_replace = msx.os.replace
    calls = {"n": 0}

    def flaky(src, dst):
        if str(dst).endswith("market_state.xlsx"):
            calls["n"] += 1
            raise PermissionError(5, "Access is denied")
        return real_replace(src, dst)

    monkeypatch.setattr(msx.os, "replace", flaky)
    monkeypatch.setattr(msx.time, "sleep", lambda *_: None)

    state = _state()
    state.update(as_of="2026-08-21", generated_at="2026-08-21T00:00:00+00:00",
                 narrative="n", conditions=ms.conditions(_state()),
                 vol_hist=pd.DataFrame(), pos_hist=pd.DataFrame(),
                 pos_latest=pd.DataFrame(), fund_hist=pd.DataFrame(),
                 liq_hist=pd.DataFrame(), stale={})
    out = msx.build_workbook(state, target)

    assert out.name == "market_state.pending.xlsx"
    assert out.exists()
    assert calls["n"] == 5                      # retried before giving up
    assert target.read_text() == "existing"     # original left untouched


# ------------------------------------------------------------ framework sheets

def test_framework_sheets_render_with_disclaimers_and_universes(tmp_path):
    """The 86 render lines had zero coverage: _state() has no `framework` key, so
    build_workbook's guard skipped both sheets entirely."""
    from openpyxl import load_workbook

    from qlib_lab import asset_frame as af
    from qlib_lab import market_state_xlsx as msx

    row = {"asset": "SPY", "price": 765.72, "inverse": False,
           "universes": "macro22,cot", "confluence": "2B / 3Br / 3ND",
           "reads": "bullish = rising price"}
    for key, _q in af.LAYERS:
        row[key] = af.NO_DATA if key in ("valuation", "fundamentals") else af.NEUTRAL
        row[f"{key}_detail"] = f"{key} detail text"
    fw = pd.DataFrame([row])

    state = _state()
    state.update(as_of="2026-08-21", generated_at="2026-08-21T00:00:00+00:00",
                 narrative="n", conditions=ms.conditions(_state()),
                 vol_hist=pd.DataFrame(), pos_hist=pd.DataFrame(),
                 pos_latest=pd.DataFrame(), fund_hist=pd.DataFrame(),
                 liq_hist=pd.DataFrame(), stale={},
                 framework=fw, framework_focus=["SPY"])

    out = msx.build_workbook(state, tmp_path / "wb.xlsx")
    wb = load_workbook(out)
    assert "Framework" in wb.sheetnames and "Framework Detail" in wb.sheetnames

    ws = wb["Framework"]
    hdr = [c.value for c in ws[8]]
    assert hdr[:4] == ["asset", "price", "reads", "universes"]
    assert "confluence" in hdr
    # the disclaimers are what keep this sheet descriptive - they must be present
    blob = " ".join(str(ws.cell(row=r, column=1).value or "") for r in range(1, 8))
    assert "COUNT" in blob and "NO_DATA" in blob
    assert "crowded" in blob.lower() and "risk-OFF" in blob
    assert ws.cell(row=9, column=4).value == "macro22,cot"


def test_framework_sheets_skipped_when_unavailable(tmp_path):
    """A framework failure must degrade the sheets, never the core workbook."""
    from openpyxl import load_workbook

    from qlib_lab import market_state_xlsx as msx
    state = _state()
    state.update(as_of="2026-08-21", generated_at="x", narrative="n",
                 conditions=ms.conditions(_state()), vol_hist=pd.DataFrame(),
                 pos_hist=pd.DataFrame(), pos_latest=pd.DataFrame(),
                 fund_hist=pd.DataFrame(), liq_hist=pd.DataFrame(), stale={},
                 framework=pd.DataFrame(), framework_focus=[])
    wb = load_workbook(msx.build_workbook(state, tmp_path / "wb2.xlsx"))
    assert "Framework" not in wb.sheetnames
    assert "TODAY" in wb.sheetnames          # core workbook still written


# ------------------------------------------------------------- state history log

def test_snapshot_records_drivers_not_just_fired_ids():
    """A log of fired IDs cannot answer 'was it close?' or 'which way is this
    drifting?' - the drivers are the point of a state record."""
    s = _state(vrp_pct=0.17, absorption=0.72)
    s["conditions"] = ms.conditions(s)
    snap = ms._snapshot(s)
    for k in ("vix", "vrp_pct", "term_ratio", "netliq_chg13w_z", "absorption",
              "btc_funding_z", "crowded_long", "crowded_short"):
        assert k in snap
    assert snap["vrp_pct"] == pytest.approx(0.17)
    assert snap["absorption"] == pytest.approx(0.72)


def test_snapshot_records_no_forward_return():
    """The tripwire: joining recorded state to future prices makes this a model and
    the PBO/DSR gate applies. Nothing forward-looking may enter the log."""
    s = _state()
    s["conditions"] = ms.conditions(s)
    snap = ms._snapshot(s)
    banned = ("fwd", "forward", "future", "ret_", "return", "pnl", "outcome",
              "target", "score", "hit", "correct")
    for key in snap:
        assert not any(b in key.lower() for b in banned), f"forward-looking key: {key}"


def test_snapshot_survives_empty_state():
    snap = ms._snapshot({"conditions": []})
    assert snap["vix"] is None and snap["crowded_long"] == []


def test_fire_log_row_carries_the_snapshot(tmp_path):
    import json
    s = _state(vrp_pct=0.17)
    s.update(as_of="2026-08-21", generated_at="2026-08-21T00:00:00+00:00")
    s["conditions"] = ms.conditions(s)
    p = tmp_path / "log.jsonl"
    ms.append_fire_log(s, p)
    row = json.loads(p.read_text().splitlines()[0])
    assert row["as_of"] == "2026-08-21"
    assert "fired" in row and "vrp_pct" in row and "absorption" in row


def test_history_sheet_renders_from_the_log(tmp_path, monkeypatch):
    import json
    from openpyxl import load_workbook

    from qlib_lab import market_state_xlsx as msx
    log = tmp_path / "log.jsonl"
    log.write_text("".join(json.dumps({
        "as_of": d, "generated_at": f"{d}T00:00:00+00:00", "fired": ["VOL_RICH"],
        "vix": 15.1, "vrp_pct": 0.17, "absorption": 0.72,
        "crowded_long": ["GOLD"], "crowded_short": [],
        "framework": {"n_assets": 89, "mean_bullish": 2.5},
    }) + "\n" for d in ("2026-08-19", "2026-08-20", "2026-08-21")), encoding="utf-8")
    monkeypatch.setattr(ms, "FIRE_LOG", log)

    state = _state()
    state.update(as_of="2026-08-21", generated_at="x", narrative="n",
                 conditions=ms.conditions(_state()), vol_hist=pd.DataFrame(),
                 pos_hist=pd.DataFrame(), pos_latest=pd.DataFrame(),
                 fund_hist=pd.DataFrame(), liq_hist=pd.DataFrame(), stale={},
                 framework=pd.DataFrame(), framework_focus=[])
    wb = load_workbook(msx.build_workbook(state, tmp_path / "wb.xlsx"))
    assert "History" in wb.sheetnames
    ws = wb["History"]
    hdr = [c.value for c in ws[6]]
    assert hdr[0] == "as_of" and "vrp_pct" in hdr and "n_fired" in hdr
    assert ws.cell(row=7, column=1).value == "2026-08-19"
    assert ws.max_row == 9                       # 3 logged sessions
    blob = str(ws.cell(row=2, column=1).value)
    assert "NO forward returns" in blob


def test_history_sheet_handles_an_empty_log(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    from qlib_lab import market_state_xlsx as msx
    monkeypatch.setattr(ms, "FIRE_LOG", tmp_path / "absent.jsonl")
    state = _state()
    state.update(as_of="x", generated_at="x", narrative="n",
                 conditions=ms.conditions(_state()), vol_hist=pd.DataFrame(),
                 pos_hist=pd.DataFrame(), pos_latest=pd.DataFrame(),
                 fund_hist=pd.DataFrame(), liq_hist=pd.DataFrame(), stale={},
                 framework=pd.DataFrame(), framework_focus=[])
    wb = load_workbook(msx.build_workbook(state, tmp_path / "wb2.xlsx"))
    assert "History" in wb.sheetnames
    assert "no history recorded yet" in str(wb["History"]["A6"].value)


# ------------------------------------------------------- registered market stories

def _story_panel(n=400, corr=None):
    """Panel with a controllable BTC/GOLD relationship; everything else is noise."""
    rng = np.random.default_rng(3)
    idx = pd.bdate_range("2024-01-01", periods=n)
    base = rng.normal(0, 0.01, n)
    out = {}
    for name in ("BTC", "GOLD", "SPY", "DXY", "COPPER", "TLT", "VIX", "HYG",
                 "EEM", "OIL", "US10Y"):
        r = rng.normal(0, 0.01, n)
        out[name] = pd.Series(100 * (1 + r).cumprod(), index=idx)
    if corr is not None:
        rb = base
        rg = corr * base + np.sqrt(max(1 - corr ** 2, 0.0)) * rng.normal(0, 1, n) * 0.01
        out["BTC"] = pd.Series(100 * (1 + rb).cumprod(), index=idx)
        out["GOLD"] = pd.Series(100 * (1 + rg).cumprod(), index=idx)
    return pd.DataFrame(out)


def test_story_state_is_signed_against_the_claims_own_direction():
    """One rule covers both signs: a story that implies NEGATIVE correlation holds when
    the measurement is negative enough, and is BROKEN when it flips positive."""
    assert ms._story_state(+0.55, +0.40) == ms.HOLDING
    assert ms._story_state(+0.25, +0.40) == ms.WEAK
    assert ms._story_state(-0.30, +0.40) == ms.BROKEN
    assert ms._story_state(-0.45, -0.30) == ms.HOLDING
    assert ms._story_state(-0.12, -0.30) == ms.WEAK
    assert ms._story_state(+0.31, -0.30) == ms.BROKEN
    assert ms._story_state(None, 0.4) == ms.NO_DATA
    assert ms._story_state(float("nan"), 0.4) == ms.NO_DATA


def test_story_checks_measure_the_registered_pair():
    holds = {s["id"]: s for s in ms.story_checks(_story_panel(corr=0.9))}
    breaks = {s["id"]: s for s in ms.story_checks(_story_panel(corr=-0.9))}
    assert holds["BTC_IS_DIGITAL_GOLD"]["state"] == ms.HOLDING
    assert breaks["BTC_IS_DIGITAL_GOLD"]["state"] == ms.BROKEN
    assert holds["BTC_IS_DIGITAL_GOLD"]["implies"] == "corr(BTC,GOLD) > +0.40"


def test_story_reports_no_data_rather_than_a_correlation_off_five_days():
    short = _story_panel(n=15)
    assert all(s["state"] == ms.NO_DATA for s in ms.story_checks(short))
    assert all(s["state"] == ms.NO_DATA for s in ms.story_checks(pd.DataFrame()))


def test_story_missing_leg_is_no_data_not_an_exception():
    panel = _story_panel().drop(columns=["GOLD"])
    st = {s["id"]: s for s in ms.story_checks(panel)}
    assert st["BTC_IS_DIGITAL_GOLD"]["state"] == ms.NO_DATA
    assert st["BTC_IS_A_RISK_ASSET"]["state"] != ms.NO_DATA     # unaffected pair


def test_stories_for_filters_to_the_assets_the_story_is_about():
    st = ms.story_checks(_story_panel())
    assert {s["id"] for s in ms.stories_for("BTC", st)} == {"BTC_IS_DIGITAL_GOLD",
                                                            "BTC_IS_A_RISK_ASSET"}
    assert ms.stories_for("XLU", st) == []


def test_snapshot_records_each_story_state_so_a_turn_gets_a_date():
    """The point of the table is the forward record: WHEN a relationship turned."""
    s = _state()
    s["conditions"] = ms.conditions(s)
    s["stories"] = ms.story_checks(_story_panel(corr=0.9))
    snap = ms._snapshot(s)
    assert snap["stories"]["BTC_IS_DIGITAL_GOLD"] == ms.HOLDING
    assert snap["story_corr"]["BTC_IS_DIGITAL_GOLD"] > 0.4
    banned = ("fwd", "forward", "future", "ret_", "return", "pnl", "outcome",
              "target", "score", "hit", "correct")
    for key in list(snap) + list(snap["stories"]) + list(snap["story_corr"]):
        assert not any(b in key.lower() for b in banned), f"forward-looking key: {key}"


def test_stories_note_keeps_the_registration_honest():
    """The set of stories was chosen with the 2026-08 tape visible. If that sentence
    ever disappears the sheet starts implying a track record it does not have."""
    assert "only rows in the log dated after registration are evidence" in ms.STORIES_NOTE
    assert "forecasts" in ms.STORIES_NOTE


# ------------------------------------------- as-of alignment (found 2026-08-27)

def test_last_row_asof_never_returns_a_row_from_after_the_header_date():
    """`_last_row` takes the newest row of each frame INDEPENDENTLY, so the sheet
    printed a netliq z and a funding z belonging to the session AFTER its own as_of.
    Same failure as the cross-date term ratio: two dates presented as one state."""
    idx = pd.to_datetime(["2026-08-21", "2026-08-24", "2026-08-25"])
    df = pd.DataFrame({"z": [-0.698522, -0.711366, -0.712514]}, index=idx)
    assert ms._last_row(df)["z"] == pytest.approx(-0.712514)          # the old behaviour
    assert ms._last_row_asof(df, "2026-08-24")["z"] == pytest.approx(-0.711366)
    assert ms._last_row_asof(df, None)["z"] == pytest.approx(-0.712514)
    # an as_of before every row falls back rather than returning nothing
    assert ms._last_row_asof(df, "2016-01-01")["z"] == pytest.approx(-0.712514)
    assert ms._last_row_asof(pd.DataFrame(), "2026-08-24") == {}


def test_absorption_from_corr_is_the_single_implementation():
    """Two copies of this eigen block drifted by 5e-4 on the same date. One copy now,
    imported by state_panel for the history."""
    c = pd.DataFrame(np.eye(4), columns=list("abcd"), index=list("abcd"))
    a, top1, n = ms.absorption_from_corr(c)
    assert n == 4 and a == pytest.approx(0.25) and top1 == pytest.approx(0.25)
    ones = pd.DataFrame(np.ones((4, 4)), columns=list("abcd"), index=list("abcd"))
    a1, top1_1, _ = ms.absorption_from_corr(ones)     # one factor explains everything
    assert a1 == pytest.approx(1.0) and top1_1 == pytest.approx(1.0)
    assert ms.absorption_from_corr(pd.DataFrame())[0] is None
